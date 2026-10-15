"""Philadelphia Fed quarterly Survey of Professional Forecasters inflation series."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import TYPE_CHECKING, ClassVar

from openpyxl import load_workbook

from opendata.data.capability import Capability
from opendata.data.models.federal_reserve_macro import FederalReserveInflationExpectations
from opendata.data.protocol import FetchContext, Fetcher
from opendata.data.providers.federal_reserve._iter2_client import (
    FederalReserveSourceError,
    get_bytes,
)
from opendata.data.providers.federal_reserve.models._iter2_common import (
    FederalReserveDateQuery,
    parse_float,
)

if TYPE_CHECKING:
    from opendata.data.models.base import ContractModel

_WORKBOOK_URL = (
    "https://www.philadelphiafed.org/-/media/FRBP/Assets/Surveys-And-Data/"
    "survey-of-professional-forecasters/historical-data/Inflation.xlsx"
    "?hash=0E7FC9E86818CF8A1C512953603AA6D8&sc_lang=en"
)
_SHEET = "INFLATION"
_HEADERS = ("YEAR", "QUARTER", "INFPGDP1YR", "INFCPI1YR", "INFCPI10YR")
_FIRST_PUBLISHED_DATE = date(1970, 4, 1)


def _quarter_date(year_value: object, quarter_value: object) -> date:
    """Convert the SPF numeric year/quarter cells to the quarter's first day."""
    if isinstance(year_value, bool) or not isinstance(year_value, (int, float)):
        raise FederalReserveSourceError("FEDERAL_RESERVE_BAD_INFLATION_PERIOD")
    if isinstance(quarter_value, bool) or not isinstance(quarter_value, (int, float)):
        raise FederalReserveSourceError("FEDERAL_RESERVE_BAD_INFLATION_PERIOD")
    if not float(year_value).is_integer() or not float(quarter_value).is_integer():
        raise FederalReserveSourceError("FEDERAL_RESERVE_BAD_INFLATION_PERIOD")
    year = int(year_value)
    quarter = int(quarter_value)
    if not 1900 <= year <= 2200 or quarter not in {1, 2, 3, 4}:
        raise FederalReserveSourceError("FEDERAL_RESERVE_BAD_INFLATION_PERIOD")
    return date(year, 1 + (quarter - 1) * 3, 1)


def _forecast_value(value: object) -> float | None:
    """Keep the SPF's explicit ``#N/A`` cells null and reject other error tokens."""
    if value is None or value == "#N/A":
        return None
    return parse_float(value, nullable=False)


class InflationExpectationsQuery(FederalReserveDateQuery):
    """Inclusive date window over SPF quarter-start observations."""


class FederalReserveInflationExpectationsFetcher(Fetcher[InflationExpectationsQuery, bytes]):
    """Parse the official SPF inflation workbook without recomputing its published measures."""

    async_mode: ClassVar[str] = "bounded_thread"
    canonical_model: ClassVar[str | None] = "InflationExpectations"
    capability: ClassVar[Capability] = Capability(
        asset_class="macro",
        domain="federal_reserve_inflation_expectations",
        period="q",
        market="us",
        source="federal_reserve",
        verified=False,
    )
    #: The contract ``domains.yaml`` declares for this domain -- verified by
    #: ``contract_model('federal_reserve_inflation_expectations') ->
    #: opendata.data.models.federal_reserve_macro.FederalReserveInflationExpectations``, which is
    #: the class :meth:`transform_data` already builds, and the governed query service revalidates
    #: every row against it rather than against this fetcher's own import.
    row_model: ClassVar[type[ContractModel]] = FederalReserveInflationExpectations

    def transform_query(self, **kwargs: object) -> InflationExpectationsQuery:
        """Validate the optional inclusive date window."""
        return InflationExpectationsQuery.model_validate(kwargs)

    def extract_data(self, params: InflationExpectationsQuery, ctx: FetchContext) -> bytes:
        """Read Philadelphia Fed's published one-year and ten-year SPF inflation workbook."""
        if type(params) is not InflationExpectationsQuery:
            raise FederalReserveSourceError("FEDERAL_RESERVE_BAD_QUERY")
        return get_bytes(_WORKBOOK_URL, ctx.timeout)

    def transform_data(
        self, raw: bytes, params: InflationExpectationsQuery
    ) -> tuple[FederalReserveInflationExpectations, ...]:
        """Require the exact published sheet/header and preserve SPF missing-value semantics."""
        if not isinstance(raw, bytes) or not raw:
            raise FederalReserveSourceError("FEDERAL_RESERVE_BAD_INFLATION_WORKBOOK")
        try:
            workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        except Exception:
            raise FederalReserveSourceError("FEDERAL_RESERVE_BAD_INFLATION_WORKBOOK") from None
        try:
            if workbook.sheetnames != [_SHEET]:
                raise FederalReserveSourceError("FEDERAL_RESERVE_BAD_INFLATION_WORKBOOK")
            worksheet = workbook[_SHEET]
            rows = worksheet.iter_rows(values_only=True)
            try:
                header = next(rows)
            except StopIteration:
                raise FederalReserveSourceError("FEDERAL_RESERVE_BAD_INFLATION_HEADER") from None
            normalized_header = tuple(
                cell.strip() if isinstance(cell, str) else cell for cell in header
            )
            if normalized_header != _HEADERS:
                raise FederalReserveSourceError("FEDERAL_RESERVE_BAD_INFLATION_HEADER")
            records: list[tuple[object, ...]] = []
            for row in rows:
                if all(cell is None for cell in row):
                    continue
                if len(row) != len(_HEADERS):
                    raise FederalReserveSourceError("FEDERAL_RESERVE_BAD_INFLATION_ROW")
                records.append(row)
        finally:
            workbook.close()

        result: list[FederalReserveInflationExpectations] = []
        seen: set[date] = set()
        for row in records:
            observation_date = _quarter_date(row[0], row[1])
            if observation_date in seen:
                raise FederalReserveSourceError("FEDERAL_RESERVE_DUPLICATE_INFLATION_DATE")
            seen.add(observation_date)
            if observation_date < _FIRST_PUBLISHED_DATE:
                continue
            if params.start_date is not None and observation_date < params.start_date:
                continue
            if params.end_date is not None and observation_date > params.end_date:
                continue
            result.append(
                FederalReserveInflationExpectations.model_validate(
                    {
                        "date": observation_date,
                        "infpgdp1yr": _forecast_value(row[2]),
                        "infcpi1yr": _forecast_value(row[3]),
                        "infcpi10yr": _forecast_value(row[4]),
                    }
                )
            )
            if len(result) > params.max_records:
                raise FederalReserveSourceError("FEDERAL_RESERVE_TOO_MANY_RECORDS")
        return tuple(result)


__all__ = ["FederalReserveInflationExpectationsFetcher", "InflationExpectationsQuery"]
