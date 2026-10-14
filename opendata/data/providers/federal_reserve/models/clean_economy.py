"""Clean-room Federal Reserve economy fetchers.

Only public New York Fed, Philadelphia Fed, and San Francisco Fed resources
are used here.  Requests go through the shared governed HTTP client and are
executed by Fetcher's bounded thread mode.
"""

from __future__ import annotations

import json
import math
import re
import time
from collections.abc import Mapping
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any, ClassVar, Literal, cast
from urllib.parse import quote

from openpyxl import load_workbook
from pydantic import ConfigDict, StrictBool, StrictStr, ValidationError, model_validator

from opendata.data.capability import Capability
from opendata.data.http_client import get_shared_http_client
from opendata.data.models.federal_reserve_clean_economy import (
    CentralBankHoldings,
    InflationExpectations,
    TotalFactorProductivity,
)
from opendata.data.protocol import FetchContext, Fetcher, QueryParams
from opendata.data.providers._engine.http_json import ProviderEngineError
from opendata.data.request_budget import (
    RequestExecutionCancelledError,
    RequestExecutionDeadlineError,
)

NYFED_BASE = "https://markets.newyorkfed.org/api/soma"
_Date = date
PHILLY_INFLATION_XLSX = (
    "https://www.philadelphiafed.org/-/media/FRBP/Assets/Surveys-And-Data/"
    "survey-of-professional-forecasters/historical-data/Inflation.xlsx"
    "?hash=0E7FC9E86818CF8A1C512953603AA6D8&sc_lang=en"
)
SF_TFP_XLSX = "https://www.frbsf.org/wp-content/uploads/quarterly_tfp.xlsx?2026-10-09="

_MAX_SOURCE_BYTES = 32 * 1024 * 1024
_HTTP_TIMEOUT = (4.0, 15.0)
_CUSIP_RE = re.compile(r"^[A-Za-z0-9]{9}$")
_QUARTER_RE = re.compile(r"^(\d{4}):Q([1-4])$")
_SUMMARY_LABELS = {
    "Full sample mean": "full_sample_mean",
    "Past 4 qtrs": "past_4_quarters",
    "Past 8 qtrs": "past_8_quarters",
    "Since 2019:4": "since_2019",
    "2004:4-2019:4": "period_2004_2019",
    "1995:4-2004:4": "period_1995_2004",
    "1973:1-1995:4": "period_1973_1995",
    "1947:1-1973:1": "period_1947_1973",
}
_TREASURY_WAM_TYPES = frozenset({"Bills", "FRNs", "NotesBonds", "TIPS"})
_AGENCY_WAM_TYPES = frozenset({"Agency Debts"})

_TFP_SOURCE_TO_FIELD = {
    "dY_prod": "d_y_prod",
    "dY_inc": "d_y_inc",
    "dY": "d_y",
    "dhours": "d_hours",
    "dLP": "d_lp",
    "dk": "d_k",
    "dLQ_BLS_interpolated": "d_lq_bls_interpolated",
    "dLQ_Aaronson_Sullivan": "d_lq_aaronson_sullivan",
    "dLQ": "d_lq",
    "alpha": "alpha",
    "dtfp": "d_tfp",
    "dutil": "d_util",
    "dtfp_util": "d_tfp_util",
    "relativePrice": "relative_price",
    "invShare": "inv_share",
    "dtfp_I": "d_tfp_i",
    "dtfp_C": "d_tfp_c",
    "du_invest": "d_u_invest",
    "du_consumption": "d_u_consumption",
    "dtfp_I_util": "d_tfp_i_util",
    "dtfp_C_util": "d_tfp_c_util",
}
_TFP_TITLES = {
    "d_y_prod": "Business output, expenditure-side growth",
    "d_y_inc": "Business output, income-side growth",
    "d_y": "Business output growth",
    "d_hours": "Business-sector hours worked growth",
    "d_lp": "Business-sector labor productivity growth",
    "d_k": "Business-sector capital input growth",
    "d_lq_bls_interpolated": "BLS labor-composition growth, interpolated before 1979",
    "d_lq_aaronson_sullivan": "Aaronson-Sullivan labor-composition growth",
    "d_lq": "Labor-composition growth used in the TFP calculation",
    "alpha": "Capital share of income",
    "d_tfp": "Business-sector total factor productivity growth",
    "d_util": "Capital and labor utilization adjustment",
    "d_tfp_util": "Utilization-adjusted total factor productivity growth",
    "relative_price": "Consumption-to-equipment relative-price growth",
    "inv_share": "Equipment and consumer-durables share of business output",
    "d_tfp_i": "Investment-goods sector total factor productivity growth",
    "d_tfp_c": "Non-equipment business output total factor productivity growth",
    "d_u_invest": "Investment-goods production utilization adjustment",
    "d_u_consumption": "Non-investment business production utilization adjustment",
    "d_tfp_i_util": "Utilization-adjusted investment-goods productivity growth",
    "d_tfp_c_util": "Utilization-adjusted non-equipment productivity growth",
}
_TFP_HEADERS = ("date", *_TFP_SOURCE_TO_FIELD)


class FederalReserveSourceError(RuntimeError):
    """A Federal Reserve endpoint or workbook violated its documented shape."""


class FederalReserveNoDataError(ProviderEngineError):
    """An exact date or requested source interval has no published observation."""

    def __init__(self, code: str) -> None:
        """Initialize the no-data error with its stable source code."""
        super().__init__(code)


class _FederalReserveQuery(QueryParams):
    """Accept and validate the service's pinned source and market identities."""

    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)
    source: Literal["federal_reserve"] = "federal_reserve"
    market: Literal["us"] = "us"
    symbol: None = None
    date: _Date | None = None

    @model_validator(mode="after")
    def _check_range(self) -> _FederalReserveQuery:
        if (
            self.start_date is not None
            and self.end_date is not None
            and self.start_date > self.end_date
        ):
            raise ValueError("start_date must not be after end_date")
        return self


class CentralBankHoldingsQuery(_FederalReserveQuery):
    """OpenBB V4 holdings selectors plus the service identity fields."""

    date: _Date | None = None
    start_date: None = None
    end_date: None = None
    holding_type: (
        Literal[
            "all_agency",
            "agency_debts",
            "mbs",
            "cmbs",
            "all_treasury",
            "bills",
            "notesbonds",
            "frn",
            "tips",
        ]
        | None
    ) = "all_treasury"
    summary: StrictBool | None = False
    cusip: StrictStr | None = None
    wam: StrictBool | None = False
    monthly: StrictBool | None = False

    @model_validator(mode="after")
    def _check_selectors(self) -> CentralBankHoldingsQuery:
        if self.cusip is not None and not _CUSIP_RE.fullmatch(self.cusip):
            raise ValueError("cusip must be a nine-character alphanumeric CUSIP")
        if self.summary is True:
            return self
        if self.wam is True:
            return self
        holding_type = self.holding_type or "all_treasury"
        if self.monthly is True and holding_type not in {
            "all_treasury",
            "bills",
            "notesbonds",
            "frn",
            "tips",
        }:
            raise ValueError("monthly is supported only for Treasury holding types")
        return self


class InflationExpectationsQuery(_FederalReserveQuery):
    """Quarterly Philly Fed SPF inflation forecast date range."""

    date: None = None


class TotalFactorProductivityQuery(_FederalReserveQuery):
    """San Francisco Fed quarterly, annual, or official-summary selector."""

    date: None = None
    frequency: Literal["quarter", "annual", "summary"] = "quarter"

    @model_validator(mode="after")
    def _check_frequency_range(self) -> TotalFactorProductivityQuery:
        if self.frequency == "summary" and (
            self.start_date is not None or self.end_date is not None
        ):
            raise ValueError("start_date/end_date apply only to quarter and annual frequency")
        return self


def _check_live(ctx: FetchContext) -> None:
    """Honor shared worker cancellation/deadline before and after transport."""
    event = ctx._thread_cancel_event
    if event is not None and event.is_set():
        raise RequestExecutionCancelledError("request execution was cancelled")
    deadline = ctx._deadline_monotonic
    if deadline is not None and time.monotonic() >= deadline:
        raise RequestExecutionDeadlineError("request execution deadline expired")


def _strict_json(response: Any) -> object:  # noqa: ANN401 - requests response seam
    """Decode ordinary JSON while refusing duplicate keys and non-finite tokens."""
    body = getattr(response, "content", None)
    if not isinstance(body, bytes) or len(body) > _MAX_SOURCE_BYTES:
        raise FederalReserveSourceError("FED_SOURCE_BODY_INVALID")

    def reject_constant(_value: str) -> None:
        raise ValueError("non-finite JSON number")

    def unique_object(pairs: list[tuple[str, object]]) -> dict[str, object]:
        result: dict[str, object] = {}
        for key, value in pairs:
            if key in result:
                raise ValueError("duplicate JSON key")
            result[key] = value
        return result

    try:
        return json.loads(
            body,
            parse_constant=reject_constant,
            object_pairs_hook=unique_object,
        )
    except (UnicodeDecodeError, ValueError, json.JSONDecodeError):
        raise FederalReserveSourceError("FED_SOURCE_JSON_INVALID") from None


def _mapping(value: object, code: str = "FED_SOURCE_SHAPE_INVALID") -> Mapping[str, object]:
    if not isinstance(value, Mapping) or any(not isinstance(key, str) for key in value):
        raise FederalReserveSourceError(code)
    return cast("Mapping[str, object]", value)


def _required_text(record: Mapping[str, object], key: str) -> str:
    if key not in record:
        raise FederalReserveSourceError("FED_SOURCE_FIELD_MISSING")
    value = record[key]
    if not isinstance(value, str) or not value.strip():
        raise FederalReserveSourceError("FED_SOURCE_TEXT_INVALID")
    return value


def _optional_text(record: Mapping[str, object], key: str) -> str | None:
    if key not in record or record[key] is None or record[key] == "":
        return None
    value = record[key]
    if not isinstance(value, str):
        raise FederalReserveSourceError("FED_SOURCE_TEXT_INVALID")
    return value


def _strict_date(value: object, *, required: bool = False) -> date | None:
    if value is None or value == "":
        if required:
            raise FederalReserveSourceError("FED_SOURCE_DATE_INVALID")
        return None
    if not isinstance(value, str):
        raise FederalReserveSourceError("FED_SOURCE_DATE_INVALID")
    try:
        parsed = date.fromisoformat(value)
    except ValueError:
        raise FederalReserveSourceError("FED_SOURCE_DATE_INVALID") from None
    if parsed.isoformat() != value:
        raise FederalReserveSourceError("FED_SOURCE_DATE_INVALID")
    return parsed


def _optional_number(value: object, *, scale: float = 1.0) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
        raise FederalReserveSourceError("FED_SOURCE_NUMBER_INVALID")
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise FederalReserveSourceError("FED_SOURCE_NUMBER_INVALID") from None
    if not number.is_finite():
        raise FederalReserveSourceError("FED_SOURCE_NUMBER_INVALID")
    result = float(number) / scale
    if not math.isfinite(result):
        raise FederalReserveSourceError("FED_SOURCE_NUMBER_INVALID")
    return result


def _records(document: object, collection: str) -> list[Mapping[str, object]]:
    top = _mapping(document)
    soma = _mapping(top.get("soma"))
    if collection not in soma:
        raise FederalReserveSourceError("FED_SOURCE_FIELD_MISSING")
    raw_rows = soma[collection]
    if not isinstance(raw_rows, list):
        raise FederalReserveSourceError("FED_SOURCE_ROWS_INVALID")
    return [_mapping(row) for row in raw_rows]


def _summary_document(document: object) -> list[Mapping[str, object]]:
    return _records(document, "summary")


def _holding_document(document: object) -> list[Mapping[str, object]]:
    return _records(document, "holdings")


def _wam_row(
    document: object,
    *,
    expected_date: date,
    allowed_security_types: frozenset[str],
    required_security_types: frozenset[str] | None = None,
) -> CentralBankHoldings:
    """Validate one official WAM bucket without changing its source scale."""
    soma = _mapping(_mapping(document).get("soma"))
    actual_date = _strict_date(soma.get("asOfDate"), required=True)
    if actual_date is None:
        raise FederalReserveSourceError("FED_SOURCE_DATE_INVALID")
    if actual_date != expected_date:
        raise FederalReserveSourceError("FED_SOMA_DATE_MISMATCH")
    if "wam" not in soma or "securityTypes" not in soma:
        raise FederalReserveSourceError("FED_SOURCE_FIELD_MISSING")
    wam = _optional_number(soma["wam"])
    raw_types = soma["securityTypes"]
    if wam is None or not isinstance(raw_types, list):
        raise FederalReserveSourceError("FED_SOMA_WAM_SHAPE_INVALID")
    security_types: list[str] = []
    for item in raw_types:
        if not isinstance(item, str) or not item:
            raise FederalReserveSourceError("FED_SOMA_WAM_SHAPE_INVALID")
        security_types.append(item)
    types_set = frozenset(security_types)
    if (
        not security_types
        or len(security_types) != len(types_set)
        or not types_set.issubset(allowed_security_types)
        or (required_security_types is not None and types_set != required_security_types)
    ):
        raise FederalReserveSourceError("FED_SOMA_WAM_SHAPE_INVALID")
    try:
        return CentralBankHoldings(
            date=actual_date,
            record_type="wam",
            wam=wam,
            security_types=security_types,
        )
    except ValidationError:
        raise FederalReserveSourceError("FED_SOURCE_ROW_CONTRACT_INVALID") from None


def _holding_selector(value: str) -> tuple[bool, str]:
    """Map a public holding selector to its official API family and route value."""
    if value in {"all_treasury", "bills", "notesbonds", "frn", "tips"}:
        return True, {
            "all_treasury": "all",
            "bills": "bills",
            "notesbonds": "notesbonds",
            "frn": "frn",
            "tips": "tips",
        }[value]
    return False, {
        "all_agency": "all",
        "agency_debts": "agency debts",
        "mbs": "mbs",
        "cmbs": "cmbs",
    }[value]


def _api_holding_type(source_type: str) -> str:
    return {
        "Bills": "bills",
        "NotesBonds": "notesbonds",
        "FRNs": "frn",
        "TIPS": "tips",
        "Agency Debts": "agency_debts",
        "MBS": "mbs",
        "CMBS": "cmbs",
    }.get(source_type, "")


def _query_holding_type(public_type: str) -> str | None:
    return {
        "bills": "Bills",
        "notesbonds": "NotesBonds",
        "frn": "FRNs",
        "tips": "TIPS",
        "agency_debts": "Agency Debts",
        "mbs": "MBS",
        "cmbs": "CMBS",
    }.get(public_type)


def _central_source_fields(record: Mapping[str, object]) -> None:
    """Validate the required field subset for each documented API record family."""
    security_type = _required_text(record, "securityType")
    _strict_date(record.get("asOfDate"), required=True)
    if "cusip" not in record:
        raise FederalReserveSourceError("FED_SOURCE_FIELD_MISSING")
    if record["cusip"] not in (None, "") and not isinstance(record["cusip"], str):
        raise FederalReserveSourceError("FED_SOURCE_TEXT_INVALID")
    if security_type in {"Bills", "NotesBonds", "FRNs", "TIPS", "Agency Debts"}:
        if "parValue" not in record:
            raise FederalReserveSourceError("FED_SOURCE_FIELD_MISSING")
    elif security_type in {"MBS", "CMBS"}:
        for field_name in ("securityDescription", "term", "currentFaceValue"):
            if field_name not in record:
                raise FederalReserveSourceError("FED_SOURCE_FIELD_MISSING")
        if security_type == "MBS" and "isAggregated" not in record:
            raise FederalReserveSourceError("FED_SOURCE_FIELD_MISSING")
    else:
        raise FederalReserveSourceError("FED_SOURCE_SECURITY_TYPE_INVALID")


def _detail_row(record: Mapping[str, object]) -> CentralBankHoldings:
    _central_source_fields(record)
    security_type = _required_text(record, "securityType")
    aggregated = _optional_text(record, "isAggregated")
    if aggregated not in (None, "Y"):
        raise FederalReserveSourceError("FED_SOURCE_AGGREGATION_FLAG_INVALID")

    par_value = None
    face_value = None
    if security_type in {"Bills", "NotesBonds", "FRNs", "TIPS", "Agency Debts"}:
        par_value = _optional_number(record.get("parValue"), scale=1000.0)
    else:
        face_value = _optional_number(record.get("currentFaceValue"), scale=1000.0)

    row = {
        "date": _strict_date(record.get("asOfDate"), required=True),
        "record_type": "holding",
        "security_type": security_type,
        "description": _optional_text(record, "securityDescription"),
        "is_aggregated": aggregated,
        "cusip": _optional_text(record, "cusip"),
        "issuer": _optional_text(record, "issuer"),
        "maturity_date": _strict_date(record.get("maturityDate")),
        "term": _optional_text(record, "term"),
        "face_value": face_value,
        "par_value": par_value,
        "coupon": _optional_number(record.get("coupon")),
        "spread": _optional_number(record.get("spread")),
        "percent_outstanding": _optional_number(record.get("percentOutstanding")),
        "change_prior_week": _optional_number(record.get("changeFromPriorWeek"), scale=1000.0),
        "change_prior_year": _optional_number(record.get("changeFromPriorYear"), scale=1000.0),
    }
    try:
        return CentralBankHoldings.model_validate(row)
    except ValidationError:
        raise FederalReserveSourceError("FED_SOURCE_ROW_CONTRACT_INVALID") from None


def _summary_row(
    record: Mapping[str, object],
    *,
    prior_total: float | None,
) -> CentralBankHoldings:
    required = (
        "asOfDate",
        "mbs",
        "cmbs",
        "tips",
        "frn",
        "tipsInflationCompensation",
        "notesbonds",
        "bills",
        "agencies",
        "total",
    )
    if any(name not in record for name in required):
        raise FederalReserveSourceError("FED_SOURCE_FIELD_MISSING")
    total = _optional_number(record["total"], scale=1000.0)
    change_week = None
    if total is not None and prior_total is not None:
        change_week = total - prior_total
    row = {
        "date": _strict_date(record["asOfDate"], required=True),
        "record_type": "summary",
        "bills": _optional_number(record["bills"], scale=1000.0),
        "frn": _optional_number(record["frn"], scale=1000.0),
        "notes_and_bonds": _optional_number(record["notesbonds"], scale=1000.0),
        "tips": _optional_number(record["tips"], scale=1000.0),
        "mbs": _optional_number(record["mbs"], scale=1000.0),
        "cmbs": _optional_number(record["cmbs"], scale=1000.0),
        "agencies": _optional_number(record["agencies"], scale=1000.0),
        "total": total,
        "tips_inflation_compensation": _optional_number(
            record["tipsInflationCompensation"], scale=1000.0
        ),
        "change_prior_week": change_week,
        # The official summary endpoint has no prior-year field.  Do not
        # derive a hidden calendar-alignment rule from the weekly series.
        "change_prior_year": None,
    }
    try:
        return CentralBankHoldings.model_validate(row)
    except ValidationError:
        raise FederalReserveSourceError("FED_SOURCE_ROW_CONTRACT_INVALID") from None


class CentralBankHoldingsFetcher(Fetcher[CentralBankHoldingsQuery, object]):
    """Read public SOMA weekly detail, monthly detail, summary, and WAM data."""

    async_mode = "bounded_thread"
    canonical_model = "CentralBankHoldings"
    capability: ClassVar[Capability] = Capability(
        asset_class="macro",
        domain="central_bank_holdings",
        period="weekly",
        market="us",
        source="federal_reserve",
        verified=False,
        notes="New York Fed SOMA public API; clean-room V4 contract mapping",
    )

    def transform_query(self, **kwargs: object) -> CentralBankHoldingsQuery:
        """Validate source identity and SOMA query selectors."""
        return CentralBankHoldingsQuery.model_validate(kwargs)

    def _get_json(self, path: str, ctx: FetchContext) -> object:
        _check_live(ctx)
        response = get_shared_http_client().get(
            f"{NYFED_BASE}/{path}",
            headers={"Accept": "application/json"},
            timeout=_HTTP_TIMEOUT,
            source=self.capability.source,
            cache_context={"model": self.canonical_model or ""},
        )
        _check_live(ctx)
        return _strict_json(response)

    def _latest_date(self, ctx: FetchContext) -> date:
        top = _mapping(self._get_json("asofdates/latest.json", ctx))
        soma = _mapping(top.get("soma"))
        values = soma.get("asOfDates")
        if not isinstance(values, list) or not values:
            raise FederalReserveNoDataError("FED_SOMA_LATEST_DATE_UNAVAILABLE")
        return cast("date", _strict_date(values[0], required=True))

    def extract_data(
        self,
        params: CentralBankHoldingsQuery,
        ctx: FetchContext,
    ) -> object:
        """Fetch the official SOMA route or the two WAM buckets."""
        if params.summary is True:
            document = self._get_json("summary.json", ctx)
            return {"mode": "summary", "document": document}

        if params.wam is True:
            exact_date = params.date or self._latest_date(ctx)
            date_path = f"asof/{exact_date.isoformat()}.json"
            treasury_document = self._get_json(f"tsy/wam/all/{date_path}", ctx)
            agency_document = self._get_json(
                f"agency/wam/agency%20debts/{date_path}",
                ctx,
            )
            return {
                "mode": "wam",
                "treasury_document": treasury_document,
                "agency_document": agency_document,
                "date": exact_date,
            }

        if params.monthly is True:
            document = self._get_json("tsy/get/monthly.json", ctx)
            return {"mode": "monthly", "document": document}

        holding_type = params.holding_type or "all_treasury"
        exact_date = params.date or self._latest_date(ctx)
        is_treasury, route = _holding_selector(holding_type)
        if params.cusip is not None:
            family = "tsy" if is_treasury else "agency"
            path = f"{family}/get/cusip/{quote(params.cusip, safe='')}.json"
        elif is_treasury:
            path = f"tsy/get/{route}/asof/{exact_date.isoformat()}.json"
        else:
            path = f"agency/get/{quote(route, safe='')}/asof/{exact_date.isoformat()}.json"
        return {
            "mode": "detail",
            "document": self._get_json(path, ctx),
            "date": exact_date,
        }

    def transform_data(
        self,
        raw: object,
        params: CentralBankHoldingsQuery,
    ) -> tuple[CentralBankHoldings, ...]:
        """Normalize official SOMA responses into typed holdings rows."""
        document = _mapping(raw)
        mode = document.get("mode")
        if mode == "summary":
            rows = _summary_document(document.get("document"))
            parsed: list[tuple[date, Mapping[str, object]]] = []
            for source_row in rows:
                parsed_date = cast("date", _strict_date(source_row.get("asOfDate"), required=True))
                parsed.append((parsed_date, source_row))
            if len({row_date for row_date, _ in parsed}) != len(parsed):
                raise FederalReserveSourceError("FED_SOURCE_DUPLICATE_DATE")
            parsed.sort(key=lambda item: item[0])
            output: list[CentralBankHoldings] = []
            previous_total: float | None = None
            found_date = False
            for row_date, source_row in parsed:
                total = _optional_number(source_row.get("total"), scale=1000.0)
                if params.date is None or row_date == params.date:
                    found_date = found_date or params.date is None or row_date == params.date
                    output.append(_summary_row(source_row, prior_total=previous_total))
                previous_total = total
            if params.date is not None and not found_date:
                raise FederalReserveNoDataError("FED_SOMA_DATE_NOT_AVAILABLE")
            return tuple(output)

        if mode == "wam":
            requested_date = document.get("date")
            if not isinstance(requested_date, date):
                raise FederalReserveSourceError("FED_SOURCE_DATE_INVALID")
            if "treasury_document" not in document or "agency_document" not in document:
                raise FederalReserveSourceError("FED_SOURCE_FIELD_MISSING")
            treasury_row = _wam_row(
                document["treasury_document"],
                expected_date=requested_date,
                allowed_security_types=_TREASURY_WAM_TYPES,
            )
            agency_row = _wam_row(
                document["agency_document"],
                expected_date=requested_date,
                allowed_security_types=_AGENCY_WAM_TYPES,
                required_security_types=_AGENCY_WAM_TYPES,
            )
            return (treasury_row, agency_row)

        if mode not in {"detail", "monthly"}:
            raise FederalReserveSourceError("FED_SOURCE_MODE_INVALID")
        source_rows = _holding_document(document.get("document"))
        output = []
        requested_date = document.get("date") if mode == "detail" else params.date
        requested_type = params.holding_type or "all_treasury"
        desired_source_type = _query_holding_type(requested_type)
        requested_cusip = params.cusip
        for source_row in source_rows:
            _central_source_fields(source_row)
            row_date = cast("date", _strict_date(source_row.get("asOfDate"), required=True))
            if requested_date is not None and row_date != requested_date:
                continue
            source_type = _required_text(source_row, "securityType")
            if desired_source_type is not None and source_type != desired_source_type:
                continue
            source_cusip = _optional_text(source_row, "cusip")
            if requested_cusip is not None and source_cusip != requested_cusip:
                continue
            output.append(_detail_row(source_row))
        if requested_date is not None and not output:
            raise FederalReserveNoDataError("FED_SOMA_DATE_NOT_AVAILABLE")
        return tuple(output)


class InflationExpectationsFetcher(Fetcher[InflationExpectationsQuery, object]):
    """Read and normalize the official Philly Fed SPF inflation workbook."""

    async_mode = "bounded_thread"
    canonical_model = "InflationExpectations"
    capability: ClassVar[Capability] = Capability(
        asset_class="macro",
        domain="inflation_expectations",
        period="quarterly",
        market="us",
        source="federal_reserve",
        verified=False,
        notes="Philadelphia Fed SPF official Inflation.xlsx; clean-room V4 contract mapping",
    )

    def transform_query(self, **kwargs: object) -> InflationExpectationsQuery:
        """Validate the service identity and optional survey date range."""
        return InflationExpectationsQuery.model_validate(kwargs)

    def extract_data(self, params: InflationExpectationsQuery, ctx: FetchContext) -> bytes:
        """Download the official Philadelphia Fed inflation workbook."""
        _check_live(ctx)
        response = get_shared_http_client().get(
            PHILLY_INFLATION_XLSX,
            headers={"Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
            timeout=_HTTP_TIMEOUT,
            source=self.capability.source,
            cache_context={"model": self.canonical_model or ""},
        )
        _check_live(ctx)
        body = getattr(response, "content", None)
        if not isinstance(body, bytes) or not body or len(body) > _MAX_SOURCE_BYTES:
            raise FederalReserveSourceError("PHILLY_SPF_WORKBOOK_INVALID")
        return body

    def transform_data(
        self,
        raw: object,
        params: InflationExpectationsQuery,
    ) -> tuple[InflationExpectations, ...]:
        """Normalize the INFLATION worksheet into typed forecast rows."""
        if not isinstance(raw, bytes):
            raise FederalReserveSourceError("PHILLY_SPF_WORKBOOK_INVALID")
        try:
            workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
            if workbook.sheetnames != ["INFLATION"]:
                raise FederalReserveSourceError("PHILLY_SPF_SHEET_INVALID")
            sheet = workbook["INFLATION"]
            iterator = sheet.iter_rows(values_only=True)
            header = next(iterator, None)
            if header != ("YEAR", "QUARTER", "INFPGDP1YR", "INFCPI1YR", "INFCPI10YR"):
                raise FederalReserveSourceError("PHILLY_SPF_HEADER_INVALID")
            output: list[InflationExpectations] = []
            source_count = 0
            for cells in iterator:
                if cells == (None, None, None, None, None):
                    continue
                if len(cells) != 5:
                    raise FederalReserveSourceError("PHILLY_SPF_ROW_INVALID")
                year = _integer_cell(cells[0], "PHILLY_SPF_YEAR_INVALID")
                quarter = _integer_cell(cells[1], "PHILLY_SPF_QUARTER_INVALID")
                if year < 1968 or not 1 <= quarter <= 4:
                    raise FederalReserveSourceError("PHILLY_SPF_PERIOD_INVALID")
                period = date(year, (quarter - 1) * 3 + 1, 1)
                source_count += 1
                # OpenBB V4 defines the series from 1970-Q2; the official
                # workbook's 1970-Q1 row is a pre-series all-#N/A placeholder.
                if period < date(1970, 4, 1):
                    continue
                if params.start_date is not None and period < params.start_date:
                    continue
                if params.end_date is not None and period > params.end_date:
                    continue
                output.append(
                    InflationExpectations(
                        date=period,
                        infpgdp1yr=_workbook_number(cells[2], "PHILLY_SPF_VALUE_INVALID"),
                        infcpi1yr=_workbook_number(cells[3], "PHILLY_SPF_VALUE_INVALID"),
                        infcpi10yr=_workbook_number(cells[4], "PHILLY_SPF_VALUE_INVALID"),
                    )
                )
            workbook.close()
        except FederalReserveSourceError:
            raise
        except Exception:
            raise FederalReserveSourceError("PHILLY_SPF_WORKBOOK_INVALID") from None
        if source_count == 0:
            raise FederalReserveSourceError("PHILLY_SPF_NO_QUARTERLY_ROWS")
        return tuple(output)


def _integer_cell(value: object, code: str) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise FederalReserveSourceError(code)
    if not math.isfinite(float(value)) or int(value) != value:
        raise FederalReserveSourceError(code)
    return int(value)


def _workbook_number(value: object, code: str) -> float | None:
    if value is None or value == "#N/A":
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise FederalReserveSourceError(code)
    numeric = float(value)
    if not math.isfinite(numeric):
        raise FederalReserveSourceError(code)
    return numeric


class TotalFactorProductivityFetcher(Fetcher[TotalFactorProductivityQuery, object]):
    """Read SF Fed quarterly/annual tabs and their explicitly published means."""

    async_mode = "bounded_thread"
    canonical_model = "TotalFactorProductivity"
    capability: ClassVar[Capability] = Capability(
        asset_class="macro",
        domain="total_factor_productivity",
        period="quarterly",
        market="us",
        source="federal_reserve",
        verified=False,
        notes="SF Fed official quarterly_tfp.xlsx; clean-room V4 contract mapping",
    )

    def transform_query(self, **kwargs: object) -> TotalFactorProductivityQuery:
        """Validate the requested TFP frequency and date range."""
        return TotalFactorProductivityQuery.model_validate(kwargs)

    def extract_data(self, params: TotalFactorProductivityQuery, ctx: FetchContext) -> bytes:
        """Download the official San Francisco Fed TFP workbook."""
        _check_live(ctx)
        response = get_shared_http_client().get(
            SF_TFP_XLSX,
            headers={"Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
            timeout=_HTTP_TIMEOUT,
            source=self.capability.source,
            cache_context={"model": self.canonical_model or "", "frequency": params.frequency},
        )
        _check_live(ctx)
        body = getattr(response, "content", None)
        if not isinstance(body, bytes) or not body or len(body) > _MAX_SOURCE_BYTES:
            raise FederalReserveSourceError("SF_TFP_WORKBOOK_INVALID")
        return body

    def transform_data(
        self,
        raw: object,
        params: TotalFactorProductivityQuery,
    ) -> tuple[TotalFactorProductivity, ...]:
        """Normalize the selected annual, quarterly, or summary worksheet."""
        if not isinstance(raw, bytes):
            raise FederalReserveSourceError("SF_TFP_WORKBOOK_INVALID")
        try:
            workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
            if not {"quarterly", "annual"}.issubset(workbook.sheetnames):
                raise FederalReserveSourceError("SF_TFP_SHEETS_INVALID")
            if params.frequency == "summary":
                output = self._summary(workbook["quarterly"])
            else:
                sheet_name = "quarterly" if params.frequency == "quarter" else "annual"
                output = self._time_series(workbook[sheet_name], params)
            workbook.close()
        except FederalReserveSourceError:
            raise
        except Exception:
            raise FederalReserveSourceError("SF_TFP_WORKBOOK_INVALID") from None
        return tuple(output)

    def _time_series(
        self,
        sheet: Any,  # noqa: ANN401 - openpyxl worksheet
        params: TotalFactorProductivityQuery,
    ) -> list[TotalFactorProductivity]:
        iterator = sheet.iter_rows(values_only=True)
        first = next(iterator, None)
        if first is None:
            raise FederalReserveSourceError("SF_TFP_SHEET_EMPTY")
        # The quarterly tab has an explanatory note before its column header;
        # the annual tab begins directly with the same header row.
        if first == _TFP_HEADERS:
            header = first
        elif isinstance(first[0], str) and first[0].startswith(
            "Note:  All variables are percent change at an annual rate"
        ):
            header = next(iterator, None)
        else:
            raise FederalReserveSourceError("SF_TFP_HEADER_INVALID")
        if header != _TFP_HEADERS:
            raise FederalReserveSourceError("SF_TFP_HEADER_INVALID")
        output: list[TotalFactorProductivity] = []
        seen_dates: set[str] = set()
        for cells in iterator:
            if not cells or all(value is None for value in cells):
                continue
            if len(cells) != len(_TFP_HEADERS):
                raise FederalReserveSourceError("SF_TFP_ROW_INVALID")
            raw_date = cells[0]
            if isinstance(raw_date, str) and raw_date in _SUMMARY_LABELS:
                continue
            parsed_date = _period_date(raw_date, quarterly=params.frequency == "quarter")
            date_text = str(raw_date)
            if date_text in seen_dates:
                raise FederalReserveSourceError("SF_TFP_DUPLICATE_DATE")
            seen_dates.add(date_text)
            if params.start_date is not None and parsed_date < params.start_date:
                continue
            if params.end_date is not None and parsed_date > params.end_date:
                continue
            fields: dict[str, object] = {"date": date_text}
            for index, source_name in enumerate(_TFP_SOURCE_TO_FIELD, start=1):
                fields[_TFP_SOURCE_TO_FIELD[source_name]] = _workbook_number(
                    cells[index], "SF_TFP_VALUE_INVALID"
                )
            try:
                output.append(TotalFactorProductivity.model_validate(fields))
            except ValidationError:
                raise FederalReserveSourceError("SF_TFP_ROW_CONTRACT_INVALID") from None
        return output

    def _summary(self, sheet: Any) -> list[TotalFactorProductivity]:  # noqa: ANN401
        iterator = sheet.iter_rows(values_only=True)
        first = next(iterator, None)
        if (
            first is None
            or not isinstance(first[0], str)
            or not first[0].startswith("Note:  All variables are percent change at an annual rate")
        ):
            raise FederalReserveSourceError("SF_TFP_HEADER_INVALID")
        header = next(iterator, None)
        if header != _TFP_HEADERS:
            raise FederalReserveSourceError("SF_TFP_HEADER_INVALID")
        summary_rows: dict[str, tuple[object, ...]] = {}
        for cells in iterator:
            if not cells or all(value is None for value in cells):
                continue
            if len(cells) != len(_TFP_HEADERS):
                raise FederalReserveSourceError("SF_TFP_ROW_INVALID")
            label = cells[0]
            if isinstance(label, str) and label in _SUMMARY_LABELS:
                if label in summary_rows:
                    raise FederalReserveSourceError("SF_TFP_DUPLICATE_SUMMARY")
                summary_rows[label] = cells
        missing = set(_SUMMARY_LABELS) - set(summary_rows)
        if missing:
            raise FederalReserveSourceError("SF_TFP_SUMMARY_ROWS_MISSING")

        output: list[TotalFactorProductivity] = []
        field_for_label = _SUMMARY_LABELS
        for index, source_name in enumerate(_TFP_SOURCE_TO_FIELD, start=1):
            variable = _TFP_SOURCE_TO_FIELD[source_name]
            fields: dict[str, object] = {
                "date": None,
                "variable": variable,
                "variable_title": _TFP_TITLES[variable],
            }
            for source_label, output_field in field_for_label.items():
                fields[output_field] = _workbook_number(
                    summary_rows[source_label][index],
                    "SF_TFP_SUMMARY_VALUE_INVALID",
                )
            try:
                output.append(TotalFactorProductivity.model_validate(fields))
            except ValidationError:
                raise FederalReserveSourceError("SF_TFP_SUMMARY_CONTRACT_INVALID") from None
        return output


def _period_date(value: object, *, quarterly: bool) -> date:
    if quarterly:
        if not isinstance(value, str):
            raise FederalReserveSourceError("SF_TFP_PERIOD_INVALID")
        match = _QUARTER_RE.fullmatch(value)
        if match is None:
            raise FederalReserveSourceError("SF_TFP_PERIOD_INVALID")
        return date(int(match.group(1)), 1 + (int(match.group(2)) - 1) * 3, 1)
    year = _integer_cell(value, "SF_TFP_PERIOD_INVALID")
    if not 1800 <= year <= 2200:
        raise FederalReserveSourceError("SF_TFP_PERIOD_INVALID")
    return date(year, 1, 1)


__all__ = [
    "CentralBankHoldingsFetcher",
    "CentralBankHoldingsQuery",
    "FederalReserveNoDataError",
    "FederalReserveSourceError",
    "InflationExpectationsFetcher",
    "InflationExpectationsQuery",
    "TotalFactorProductivityFetcher",
    "TotalFactorProductivityQuery",
]
