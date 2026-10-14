"""Offline EIA WPSR and STEO fixtures, limits, and service-path checks."""

from __future__ import annotations

import asyncio
import json
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace
from typing import Any

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from opendata.data.models.eia import PetroleumStatusReport, ShortTermEnergyOutlook
from opendata.data.protocol import FetchContext
from opendata.data.providers.eia import provider as eia_provider
from opendata.data.providers.eia.models import (
    _client,
    short_term_energy_outlook,
)
from opendata.data.providers.eia.models.petroleum_status_report import (
    PetroleumStatusReportFetcher,
)
from opendata.data.providers.eia.models.short_term_energy_outlook import (
    ShortTermEnergyOutlookFetcher,
)
from opendata.data.request_budget import (
    GrantDecision,
    RequestBudget,
    RequestBudgetProfile,
    RequestGrant,
    RequestOperation,
)
from opendata.services import provider_model_query


class _Response:
    def __init__(
        self,
        payload: object,
        *,
        status_code: int = 200,
        content_type: str = "application/json",
    ) -> None:
        self.status_code = status_code
        if isinstance(payload, bytes):
            self.content = payload
        else:
            self.content = json.dumps(payload).encode("utf-8")
        self.headers = {"Content-Type": content_type}


class _Transport:
    def __init__(self, *responses: _Response) -> None:
        self.responses = list(responses)
        self.requests: list[dict[str, Any]] = []

    def get(
        self,
        url: str,
        *,
        params: dict[str, str] | None = None,
        source: str,
        **kwargs: object,
    ) -> _Response:
        self.requests.append({"url": url, "params": dict(params or {}), "source": source})
        if not self.responses:
            raise AssertionError("unexpected extra EIA request")
        return self.responses.pop(0)


class _Principal:
    def allows_domain(self, domain: str) -> bool:
        return domain.startswith("energy_")


def _context(model: str) -> FetchContext:
    grant = RequestGrant(
        source="eia",
        canonical_model=model,
        operation=RequestOperation.QUERY,
        decision=GrantDecision.ALLOWED,
        rights_evidence="test-fixture:offline-only:fake-transport",
        task_attempts=10,
        source_attempts=10,
        allowed_hosts=frozenset({"api.eia.gov", "ir.eia.gov", "www.eia.gov"}),
        expires_at=datetime.now(timezone.utc) + timedelta(hours=1),
        profile=RequestBudgetProfile.BOUNDED_REQUEST,
    )
    budget = RequestBudget(
        task_attempts=10,
        source_attempts=10,
        grants=(grant,),
        profile=RequestBudgetProfile.BOUNDED_REQUEST,
    )
    return FetchContext(timeout=2.0, request_budget=budget)


def _install_transport(monkeypatch: pytest.MonkeyPatch, *responses: _Response) -> _Transport:
    transport = _Transport(*responses)
    monkeypatch.setattr(_client, "get_shared_http_client", lambda: transport)
    return transport


def _descriptor(fetcher: object, model: str) -> SimpleNamespace:
    capability = fetcher.capability  # type: ignore[attr-defined]
    return SimpleNamespace(
        source=capability.source,
        model=model,
        domain=capability.domain,
        verified=capability.verified,
        full_capability_identity=(
            capability.asset_class,
            capability.domain,
            capability.period,
            capability.market,
            capability.source,
        ),
    )


def _install_service_path(
    monkeypatch: pytest.MonkeyPatch,
    fetcher: object,
    model: str,
    contract: type,
) -> None:
    monkeypatch.setattr(
        provider_model_query,
        "_resolve_fetcher",
        lambda *_args: (_descriptor(fetcher, model), fetcher),
    )
    monkeypatch.setattr(
        provider_model_query,
        "_reviewed_query_contract",
        lambda _domain: (None, contract),
    )


def _steo_page(data: list[dict[str, object]], total: int = 1) -> _Response:
    return _Response(
        {
            "response": {
                "total": str(total),
                "frequency": "monthly",
                "dateFormat": "YYYY-MM",
                "data": data,
            }
        }
    )


_WPSR_CSV = b"""STUB_1,STUB_2,STUB_3,10/09/2026,10/02/2026
U.S. Petroleum Balance Sheet,,,
Supply,Production,Crude Oil,13.4,13.2
Supply,Imports,Total,6.1,6.2
Ending Stocks,Crude Oil,Total,420.3,421.0
"""


def _steo_archive_page() -> bytes:
    return b"""<!doctype html><html><body><table>
    <tr><td>October 2026</td><td>10/06/2026</td>
    <td><a href="archives/oct26.pdf">oct26.pdf</a></td>
    <td><a href="archives/oct26_base.xlsx">oct26_base.xlsx</a></td></tr>
    <tr><td>September 2026</td><td>09/09/2026</td>
    <td><a href="archives/sep26.pdf">sep26.pdf</a></td>
    <td><a href="archives/sep26_base.xlsx">sep26_base.xlsx</a></td></tr>
    </table></body></html>"""


def _steo_archive_workbook(
    *,
    crude_values: tuple[float, float] = (11.4, 11.5),
    oil_values: tuple[float, float] = (80.0, 81.0),
    malformed_marker: bool = False,
) -> bytes:
    workbook = Workbook()
    dates = workbook.active
    dates.title = "Dates"
    dates["A1"] = "Forecast Month -"
    dates["D1"] = "October 2026"
    dates["A2"] = "Modeling and analysis completion -"
    dates["D2"] = datetime(2026, 10, 1)
    dates["A7"] = "Last Historical Month---"
    dates["D7"] = 202609
    dates["B11"] = "AAAA_DATEX or AAAA_YEAR"
    dates["B12"] = "Column ID"
    dates["B13"] = "Historical"
    for month in range(1, 13):
        column = month + 2
        dates.cell(11, column, 202600 + month)
        dates.cell(12, column, 300 + month)
        dates.cell(13, column, "invalid" if malformed_marker and month == 10 else int(month <= 9))

    contents = workbook.create_sheet("Contents")
    contents["A8"] = "Table 1. U.S. Energy Markets Summary"
    contents["A9"] = "Table 2. Nominal Energy Prices"

    table_one = workbook.create_sheet("1tab")
    table_one["B1"] = "Table 1. U.S. Energy Markets Summary"
    table_one["A3"] = "Forecast date:"
    table_one["C3"] = 2026
    table_one["A4"] = "Thursday, October 1, 2026"
    for month in range(1, 13):
        table_one.cell(4, month + 2, date(2026, month, 1).strftime("%b"))
    table_one["B5"] = "Energy Production"
    table_one["B7"] = "Crude Oil Production (a)"
    table_one["A8"] = "COPRPUS"
    table_one["B8"] = "(million barrels per day)"
    table_one.cell(8, 11, crude_values[0])
    table_one.cell(8, 12, crude_values[1])
    table_one["B10"] = "Natural Gas"
    table_one["A11"] = "NGPRPUS"
    table_one["B11"] = "(billion cubic feet per day)"
    table_one.cell(11, 3, 95.1)
    table_one.cell(11, 4, 96.2)
    table_one["B13"] = "Weather"
    table_one["A14"] = "ZWHDPUS"
    table_one["B14"] = "U.S. Heating Degree-Days"
    table_one.cell(14, 11, 100.0)
    table_one.cell(14, 12, 50.0)

    table_two = workbook.create_sheet("2tab")
    table_two["B1"] = "Table 2. Nominal Energy Prices"
    table_two["A3"] = "Forecast date:"
    table_two["C3"] = 2026
    table_two["A4"] = "Thursday, October 1, 2026"
    for month in range(1, 13):
        table_two.cell(4, month + 2, date(2026, month, 1).strftime("%b"))
    table_two["B5"] = "Crude Oil (dollars per barrel)"
    table_two["A6"] = "WTIPUUS"
    table_two["B6"] = "West Texas Intermediate Spot Average"
    table_two.cell(6, 11, oil_values[0])
    table_two.cell(6, 12, oil_values[1])
    table_two["B7"] = "Gasoline (dollars per gallon)"
    table_two["A8"] = "MGEIAUS_$"
    table_two["B8"] = "All grades United States average"
    table_two.cell(8, 11, 3.5)
    table_two.cell(8, 12, 3.6)

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_psr_service_injects_identity_and_serializes_official_csv_shape(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.delenv("EIA_API_KEY", raising=False)
    transport = _install_transport(
        monkeypatch,
        _Response(_WPSR_CSV, content_type="text/csv"),
    )
    fetcher = PetroleumStatusReportFetcher()
    _install_service_path(monkeypatch, fetcher, "PetroleumStatusReport", PetroleumStatusReport)
    received_query: dict[str, object] = {}
    original_transform = fetcher.transform_query

    def capture_transform(**kwargs: object) -> object:
        received_query.update(kwargs)
        return original_transform(**kwargs)

    monkeypatch.setattr(fetcher, "transform_query", capture_transform)
    response = asyncio.run(
        provider_model_query.query_provider_model(
            registry=object(),
            source="eia",
            model="PetroleumStatusReport",
            query={
                "category": "balance_sheet",
                "table": "crude",
                "start_date": "2026-10-09",
            },
            principal=_Principal(),
            ctx=_context("PetroleumStatusReport"),
        )
    )

    assert received_query["source"] == "eia"
    assert received_query["market"] == "us"
    assert response["source"] == "eia"
    assert response["model"] == "PetroleumStatusReport"
    assert response["results"] == [
        {
            "date": "2026-10-09",
            "table": "U.S. Petroleum Balance Sheet",
            "symbol": "table1:row:Supply/Production/Crude%20Oil",
            "order": 2,
            "title": "Supply / Production / Crude Oil",
            "value": 13.4,
            "unit": None,
        },
        {
            "date": "2026-10-09",
            "table": "U.S. Petroleum Balance Sheet",
            "symbol": "table1:row:Ending%20Stocks/Crude%20Oil/Total",
            "order": 4,
            "title": "Ending Stocks / Crude Oil / Total",
            "value": 420.3,
            "unit": None,
        },
    ]
    assert transport.requests[0]["url"] == "https://ir.eia.gov/wpsr/table1.csv"
    assert transport.requests[0]["source"] == "eia"


def test_psr_missing_values_are_omitted_and_malformed_table_fails_closed(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fetcher = PetroleumStatusReportFetcher()
    query = fetcher.transform_query(source="eia", market="us")
    empty_csv = b"STUB_1,STUB_2,10/09/2026\nSupply,Production,NA\n"
    assert fetcher.transform_data(("table1.csv", empty_csv), query) == ()
    with pytest.raises(_client.EIAProviderError, match="EIA_BAD_CSV_SHAPE"):
        fetcher.transform_data(("table1.csv", b"Date,Value\n2026-10-09,1\n"), query)
    malformed_numeric = b"STUB_1,STUB_2,10/09/2026\nSupply,Crude Oil,1x\n"
    with pytest.raises(_client.EIAProviderError, match="EIA_BAD_NUMBER"):
        fetcher.transform_data(("table1.csv", malformed_numeric), query)
    transport = _install_transport(monkeypatch)
    with pytest.raises(ValidationError):
        fetcher.transform_query(table="not-a-wpsr-table")
    assert transport.requests == []


def test_psr_table_defaults_follow_public_category_contract() -> None:
    fetcher = PetroleumStatusReportFetcher()
    assert fetcher.transform_query().table == "all"
    assert fetcher.transform_query(category="weekly_estimates").table == "stocks"
    assert fetcher.transform_query(category="weekly_estimates", table=None).table == "stocks"
    assert fetcher.transform_query(category="imports", table=None).table == "all"
    assert fetcher.transform_query(category="balance_sheet", table=["all"]).table == "all"


def test_psr_only_accepts_category_table_pairs_with_public_shape_evidence(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fetcher = PetroleumStatusReportFetcher()
    transport = _install_transport(monkeypatch)

    assert fetcher.transform_query(category="balance_sheet", table="crude").table == "crude"
    assert fetcher.transform_query(category="weekly_estimates", table="imports").table == "imports"
    assert fetcher.transform_query(
        category="weekly_estimates", table=["stocks", "imports"]
    ).table == ["stocks", "imports"]

    for category, table in (
        ("balance_sheet", "monthly"),
        ("weekly_estimates", "all"),
        ("weekly_estimates", "monthly"),
        ("crude_petroleum_stocks", "monthly"),
    ):
        with pytest.raises(ValidationError, match="not verified for the selected WPSR category"):
            fetcher.transform_query(category=category, table=table)

    assert transport.requests == []


def test_psr_incompatible_category_table_fails_before_service_transport(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.delenv("EIA_API_KEY", raising=False)
    transport = _install_transport(monkeypatch)
    fetcher = PetroleumStatusReportFetcher()
    _install_service_path(monkeypatch, fetcher, "PetroleumStatusReport", PetroleumStatusReport)

    with pytest.raises(provider_model_query.ProviderModelQueryValidationError):
        asyncio.run(
            provider_model_query.query_provider_model(
                registry=object(),
                source="eia",
                model="PetroleumStatusReport",
                query={"category": "balance_sheet", "table": "monthly"},
                principal=_Principal(),
                ctx=_context("PetroleumStatusReport"),
            )
        )

    assert transport.requests == []


def test_psr_locator_uses_reversible_stub_path_and_survives_row_reordering() -> None:
    fetcher = PetroleumStatusReportFetcher()
    query = fetcher.transform_query(source="eia", market="us", table="all")
    first = (
        b"STUB_1,STUB_2,STUB_3,10/09/2026\n"
        b"Supply,Production,Crude Oil,13.4\n"
        b"Ending Stocks,Crude Oil,Total,420.3\n"
    )
    reversed_rows = (
        b"STUB_1,STUB_2,STUB_3,10/09/2026\n"
        b"Ending Stocks,Crude Oil,Total,420.3\n"
        b"Supply,Production,Crude Oil,13.4\n"
    )

    first_rows = fetcher.transform_data(("table1.csv", first), query)
    reordered_rows = fetcher.transform_data(("table1.csv", reversed_rows), query)

    assert {row.value: row.symbol for row in first_rows} == {
        13.4: "table1:row:Supply/Production/Crude%20Oil",
        420.3: "table1:row:Ending%20Stocks/Crude%20Oil/Total",
    }
    assert {row.value: row.symbol for row in reordered_rows} == {
        13.4: "table1:row:Supply/Production/Crude%20Oil",
        420.3: "table1:row:Ending%20Stocks/Crude%20Oil/Total",
    }


def test_steo_service_injects_identity_serializes_and_pages(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("EIA_API_KEY", "fixture-only-key")
    monkeypatch.setattr(short_term_energy_outlook, "PAGE_SIZE", 2)
    transport = _install_transport(
        monkeypatch,
        _steo_page(
            [
                {
                    "period": "2026-01",
                    "seriesId": "PET.WCESTUS1.M",
                    "seriesName": "U.S. Ending Stocks of Crude Oil",
                    "value": "420.3",
                    "value-units": "million barrels",
                },
                {
                    "period": "2026-02",
                    "seriesId": "PET.WCESTUS1.M",
                    "seriesName": "U.S. Ending Stocks of Crude Oil",
                    "value": "421.0",
                    "value-units": "million barrels",
                },
            ],
            total=3,
        ),
        _steo_page(
            [
                {
                    "period": "2026-03",
                    "seriesId": "PET.WCESTUS1.M",
                    "seriesName": "U.S. Ending Stocks of Crude Oil",
                    "value": "422.5",
                    "value-units": "million barrels",
                }
            ],
            total=3,
        ),
    )
    fetcher = ShortTermEnergyOutlookFetcher()
    _install_service_path(monkeypatch, fetcher, "ShortTermEnergyOutlook", ShortTermEnergyOutlook)
    received_query: dict[str, object] = {}
    original_transform = fetcher.transform_query

    def capture_transform(**kwargs: object) -> object:
        received_query.update(kwargs)
        return original_transform(**kwargs)

    monkeypatch.setattr(fetcher, "transform_query", capture_transform)
    response = asyncio.run(
        provider_model_query.query_provider_model(
            registry=object(),
            source="eia",
            model="ShortTermEnergyOutlook",
            query={"symbol": "PET.WCESTUS1.M", "start_date": "2026-02-01"},
            principal=_Principal(),
            ctx=_context("ShortTermEnergyOutlook"),
        )
    )

    assert received_query["source"] == "eia"
    assert received_query["market"] == "us"
    assert response["results"] == [
        {
            "date": "2026-02-01",
            "release_date": None,
            "is_forecast": None,
            "table": None,
            "symbol": "PET.WCESTUS1.M",
            "order": None,
            "title": "U.S. Ending Stocks of Crude Oil",
            "value": 421.0,
            "unit": "million barrels",
        },
        {
            "date": "2026-03-01",
            "release_date": None,
            "is_forecast": None,
            "table": None,
            "symbol": "PET.WCESTUS1.M",
            "order": None,
            "title": "U.S. Ending Stocks of Crude Oil",
            "value": 422.5,
            "unit": "million barrels",
        },
    ]
    assert len(transport.requests) == 2
    assert transport.requests[0]["params"]["api_key"] == "fixture-only-key"
    assert transport.requests[0]["params"]["offset"] == "0"
    assert transport.requests[0]["params"]["length"] == "2"
    assert transport.requests[1]["params"]["offset"] == "2"


def test_steo_table_selectors_and_api_symbol_override_are_validated_before_io(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    transport = _install_transport(monkeypatch)
    fetcher = ShortTermEnergyOutlookFetcher()
    assert fetcher.transform_query().table == "01"
    assert fetcher.transform_query(table=None).table == "01"
    with pytest.raises(ValidationError):
        fetcher.transform_query(table="99", symbol="PET.WCESTUS1.M")
    with pytest.raises(ValidationError):
        fetcher.transform_query(frequency="annual")
    with pytest.raises(ValidationError):
        fetcher.transform_query(release_date="2026-10-06T00:00:00")
    # The public reference says an explicit symbol overrides table selection.
    assert fetcher.transform_query(symbol="PET.WCESTUS1.M", table="01").symbol == "PET.WCESTUS1.M"
    with pytest.raises(ValidationError):
        fetcher.transform_query(symbol="bad&api_key=oops")
    assert fetcher.transform_query(symbol="MGEIAUS_$", release_date="2026-10-06").symbol == (
        "MGEIAUS_$"
    )
    assert transport.requests == []


def test_steo_default_table_uses_public_archive_and_serializes_release_markers(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.delenv("EIA_API_KEY", raising=False)
    archive = _steo_archive_workbook()
    transport = _install_transport(
        monkeypatch,
        _Response(_steo_archive_page(), content_type="text/html"),
        _Response(
            archive,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
    )
    fetcher = ShortTermEnergyOutlookFetcher()
    _install_service_path(monkeypatch, fetcher, "ShortTermEnergyOutlook", ShortTermEnergyOutlook)
    received_query: dict[str, object] = {}
    original_transform = fetcher.transform_query

    def capture_transform(**kwargs: object) -> object:
        received_query.update(kwargs)
        return original_transform(**kwargs)

    monkeypatch.setattr(fetcher, "transform_query", capture_transform)
    response = asyncio.run(
        provider_model_query.query_provider_model(
            registry=object(),
            source="eia",
            model="ShortTermEnergyOutlook",
            query={"start_date": "2026-09-01", "end_date": "2026-10-01"},
            principal=_Principal(),
            ctx=_context("ShortTermEnergyOutlook"),
        )
    )

    assert received_query["source"] == "eia"
    assert received_query["market"] == "us"
    assert response["results"] == [
        {
            "date": "2026-09-01",
            "release_date": "2026-10-06",
            "is_forecast": False,
            "table": "01",
            "symbol": "COPRPUS",
            "order": 1,
            "title": "Crude Oil Production (a)",
            "value": 11.4,
            "unit": "million barrels per day",
        },
        {
            "date": "2026-10-01",
            "release_date": "2026-10-06",
            "is_forecast": True,
            "table": "01",
            "symbol": "COPRPUS",
            "order": 1,
            "title": "Crude Oil Production (a)",
            "value": 11.5,
            "unit": "million barrels per day",
        },
        {
            "date": "2026-09-01",
            "release_date": "2026-10-06",
            "is_forecast": False,
            "table": "01",
            "symbol": "ZWHDPUS",
            "order": 3,
            "title": "U.S. Heating Degree-Days",
            "value": 100.0,
            "unit": None,
        },
        {
            "date": "2026-10-01",
            "release_date": "2026-10-06",
            "is_forecast": True,
            "table": "01",
            "symbol": "ZWHDPUS",
            "order": 3,
            "title": "U.S. Heating Degree-Days",
            "value": 50.0,
            "unit": None,
        },
    ]
    assert [request["url"] for request in transport.requests] == [
        short_term_energy_outlook.ARCHIVE_URL,
        "https://www.eia.gov/outlooks/steo/archives/oct26_base.xlsx",
    ]
    assert all(request["source"] == "eia" for request in transport.requests)
    assert all(request["params"] == {} for request in transport.requests)


def test_steo_archived_release_and_nondefault_table_are_selected_exactly(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    older_release = _steo_archive_workbook(oil_values=(74.0, 75.0))
    transport = _install_transport(
        monkeypatch,
        _Response(_steo_archive_page(), content_type="text/html"),
        _Response(
            older_release,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
    )
    rows = ShortTermEnergyOutlookFetcher().fetch(
        ctx=_context("ShortTermEnergyOutlook"),
        source="eia",
        market="us",
        table="02",
        release_date="2026-09-09",
        start_date="2026-09-01",
        end_date="2026-09-01",
    )
    assert len(rows) == 2
    assert rows[0].date == date(2026, 9, 1)
    assert rows[0].release_date == date(2026, 9, 9)
    assert rows[0].is_forecast is False
    assert rows[0].table == "02"
    assert rows[0].symbol == "WTIPUUS"
    assert rows[0].value == 74.0
    assert rows[0].unit == "dollars per barrel"
    assert rows[1].symbol == "MGEIAUS_$"
    assert rows[1].value == 3.5
    assert rows[1].unit == "dollars per gallon"
    assert transport.requests[1]["url"].endswith("sep26_base.xlsx")


def test_steo_archive_rejects_missing_release_and_bad_forecast_markers(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    transport = _install_transport(
        monkeypatch,
        _Response(_steo_archive_page(), content_type="text/html"),
    )
    with pytest.raises(_client.EIAProviderError, match="EIA_STEO_RELEASE_NOT_FOUND"):
        ShortTermEnergyOutlookFetcher().fetch(
            ctx=_context("ShortTermEnergyOutlook"),
            source="eia",
            market="us",
            release_date="2026-09-08",
        )
    assert len(transport.requests) == 1

    malformed = _steo_archive_workbook(malformed_marker=True)
    transport = _install_transport(
        monkeypatch,
        _Response(_steo_archive_page(), content_type="text/html"),
        _Response(
            malformed,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
    )
    with pytest.raises(
        _client.EIAProviderError,
        match="EIA_STEO_HISTORICAL_MARKER_INVALID",
    ):
        ShortTermEnergyOutlookFetcher().fetch(
            ctx=_context("ShortTermEnergyOutlook"),
            source="eia",
            market="us",
        )
    assert len(transport.requests) == 2


def test_steo_requires_key_before_io_and_omits_missing_nonnullable_observations(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.delenv("EIA_API_KEY", raising=False)
    transport = _install_transport(monkeypatch)
    with pytest.raises(_client.EIAProviderError, match="EIA_API_KEY_MISSING"):
        ShortTermEnergyOutlookFetcher().fetch(
            ctx=_context("ShortTermEnergyOutlook"),
            source="eia",
            market="us",
            symbol="PET.WCESTUS1.M",
        )
    assert transport.requests == []

    fetcher = ShortTermEnergyOutlookFetcher()
    query = fetcher.transform_query(source="eia", market="us", symbol="PET.WCESTUS1.M")
    assert fetcher.transform_data([], query) == ()
    source_identity = {"period": "2026-01", "seriesId": "PET.WCESTUS1.M"}
    for missing_value in (
        {**source_identity, "value": None},
        {**source_identity, "value": ""},
        {**source_identity, "value": "W"},
        source_identity,
    ):
        assert fetcher.transform_data([missing_value], query) == ()
    mixed_rows = fetcher.transform_data(
        [
            {**source_identity, "value": None},
            {**source_identity, "value": "421.0", "value-units": "million barrels"},
        ],
        query,
    )
    assert len(mixed_rows) == 1
    assert mixed_rows[0].value == 421.0
    with pytest.raises(_client.EIAProviderError, match="EIA_BAD_NUMBER"):
        fetcher.transform_data(
            [
                {
                    "period": "2026-01",
                    "seriesId": "PET.WCESTUS1.M",
                    "value": "not-a-number",
                }
            ],
            query,
        )


def test_eia_rate_limit_is_reported_and_public_psr_needs_no_key(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.delenv("EIA_API_KEY", raising=False)
    transport = _Transport(_Response({}, status_code=429))
    monkeypatch.setattr(_client, "get_shared_http_client", lambda: transport)
    with pytest.raises(_client.EIAProviderError, match="EIA_HTTP_ERROR status=429"):
        PetroleumStatusReportFetcher().fetch(
            ctx=_context("PetroleumStatusReport"),
            source="eia",
            market="us",
        )
    assert transport.requests[0]["url"].endswith("table1.csv")
    assert eia_provider.PROVIDER.health_check().status == "ready"
